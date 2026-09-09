"""
Extrae ingredientes y recetas del Excel de costos a un JSON intermedio.

Uso:
    python scripts/migracion/extract_excel.py

Entrada:  docs/Costos Pattiserie Septiembre 2024.xlsx
Salida:   scripts/migracion/_extract.json
"""
import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
XLSX = ROOT / "docs" / "Costos Pattiserie Septiembre 2024.xlsx"
OUT = Path(__file__).resolve().parent / "_extract.json"

FOOTER_KEYS = {
    "cantidad de masitas",
    "cantidad de porciones",
    "costo receta",
    "costo docena",
    "costo individual",
    "cantidad de recetas:",
}


def main() -> None:
    if not XLSX.exists():
        sys.exit(f"No se encuentra el Excel: {XLSX}")

    wb = openpyxl.load_workbook(XLSX, data_only=True)

    # ── Lista maestra de insumos: hoja "Valores y Totales" ──
    ws = wb["Valores y Totales"]
    ingredients = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = row[0]
        if not name or not isinstance(name, str):
            continue
        ingredients.append(
            {
                "name": name.strip(),
                "per_kg": row[1],
                "per_gram": row[2],
            }
        )

    # ── Hojas de receta ──
    recipes = {}
    for sname in wb.sheetnames:
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        hdr = rows[0]
        is_recipe = len(hdr) > 3 and hdr[1] == "Ingrediente" and hdr[2] == "Gramos"
        if not is_recipe:
            continue

        ings = []
        footer = {}
        for r in rows[1:]:
            label = r[1] or ""
            if isinstance(label, str) and label.strip().lower() in FOOTER_KEYS:
                footer[label.strip()] = {
                    "una_receta": r[2],
                    "por_cantidad": r[3],
                    "nota": r[4] if len(r) > 4 else None,
                }
                continue
            if isinstance(label, str) and label.strip() and label.strip() != "Ingrediente":
                gramos = r[2]
                precio = r[3]
                if gramos in (None, "", 0) and precio in (None, "", 0):
                    continue
                ings.append({"ing": label.strip(), "qty": gramos, "line_cost": precio})

        recipes[sname] = {"ingredients": ings, "footer": footer}

    OUT.write_text(
        json.dumps({"ingredients": ingredients, "recipes": recipes}, ensure_ascii=False, indent=1),
        encoding="utf-8",
    )
    print(f"Insumos: {len(ingredients)}")
    print(f"Hojas de receta: {len(recipes)}")
    print(f"Escrito: {OUT}")


if __name__ == "__main__":
    main()
